#!/usr/bin/env python3
"""
Skrip Otomasi Pengambilan Tangkapan Layar Slide HTML dan Konversi ke PPTX
Memanfaatkan Playwright (Chromium/Google Chrome) dan python-pptx.
"""

import asyncio
import os
import shutil
from pathlib import Path
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

BASE_DIR = Path(__file__).resolve().parent
HTML_FILE = BASE_DIR / "index.html"
SCREENSHOTS_DIR = BASE_DIR / "screenshots"
OUTPUT_PPTX_NAME = BASE_DIR / "presentasi_analisis_faktor_risiko_serangan_jantung.pptx"
SHORT_PPTX_NAME = BASE_DIR / "presentasi.pptx"

TOTAL_SLIDES = 19
VIEWPORT_WIDTH = 1920
VIEWPORT_HEIGHT = 1080


async def capture_slides():
    SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, executable_path="/usr/bin/google-chrome")
        page = await browser.new_page(viewport={"width": VIEWPORT_WIDTH, "height": VIEWPORT_HEIGHT})
        
        await page.goto(f"file://{HTML_FILE}", wait_until="networkidle")
        await page.wait_for_timeout(800)

        # Sembunyikan elemen navigasi interaktif web agar slide tampak bersih
        await page.add_style_tag(content="""
            .header-actions, .keyboard-hint, .footer-nav {
                display: none !important;
            }
        """)

        for i in range(TOTAL_SLIDES):
            img_path = SCREENSHOTS_DIR / f"slide_{i+1:02d}.png"
            # Beri jeda agar animasi transisi dan rendering Chart.js selesai
            await page.wait_for_timeout(600)
            await page.screenshot(path=str(img_path))
            print(f"Berhasil mengambil tangkapan layar slide {i+1:02d}")
            if i < TOTAL_SLIDES - 1:
                await page.keyboard.press("ArrowRight")

        await browser.close()


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slide_images = sorted([f for f in SCREENSHOTS_DIR.iterdir() if f.name.startswith("slide_") and f.suffix == ".png"])

    for img_path in slide_images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height)

    prs.save(str(OUTPUT_PPTX_NAME))
    shutil.copyfile(str(OUTPUT_PPTX_NAME), str(SHORT_PPTX_NAME))
    print(f"Berkas PPTX berhasil dibuat: {OUTPUT_PPTX_NAME}")


if __name__ == "__main__":
    asyncio.run(capture_slides())
    build_pptx()
